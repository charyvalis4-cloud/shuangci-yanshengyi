# -*- coding: utf-8 -*-
import html, os, random, sys, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C

L = "ABCDEFGH"
esc = lambda s: html.escape(s, quote=True)
TITLE = C.TITLE
WORDS = {w['w']: w for w in C.W}


def group_of(word):
    g = set()
    for s in C.CONFUSE:
        if word in s:
            g |= s
    g.add(word)
    return g


def explain(w):
    d = WORDS[w]
    p = "；".join("%s，%s" % (c, t) for c, t in d['parts'])
    return "%s：%s。%s「%s」。" % (w, p, d['kind'], d['d'].split('。')[0])


def shuffled_mc(stem, correct, wrongs, rng, x=""):
    opts = [correct] + list(wrongs)
    rng.shuffle(opts)
    a = opts.index(correct)
    return dict(kind='mc', stem=stem, opts=opts, a=L[a], ans="%s %s" % (L[a], correct) if len(correct) < 14 else L[a], x=x)


def build_sections():
    secs = []
    # 一、字形
    qs = [dict(kind='in', stem=s, a=a, ans=a) for s, a in C.GLYPH]
    secs.append(("一、字形測驗（填入正確的字）", "", qs))
    # 二、綜合測驗
    rng = random.Random(11)
    qs = [shuffled_mc(s, c, w, rng, x) for s, c, w, x in C.MC]
    secs.append(("二、綜合測驗", "", qs))
    # 三、連連看
    qs = []
    for i, w in enumerate(C.MATCH_ITEMS):
        qs.append(dict(kind='sel', stem=w + " ⇄", opts=C.MATCH_OPTS, a=L[C.MATCH_ANS[i]],
                       ans="%s %s" % (L[C.MATCH_ANS[i]], C.MATCH_OPTS[C.MATCH_ANS[i]]), x=explain(w)))
    secs.append(("三、連連看（把詞語和引申義配起來）", "", qs))
    # 四、詞語運用
    rng = random.Random(23)
    qs = []
    allw = [x['w'] for x in C.W]
    for d in C.W:
        w = d['w']
        stem = d['ex'].replace(w, "＿＿", 1)
        pool = [x for x in allw if x not in group_of(w)]
        wr = rng.sample(pool, 3)
        qs.append(shuffled_mc(stem, w, wr, rng, explain(w)))
    secs.append(("四、詞語運用（選出最適合的詞語）", "", qs))
    # 五、本義線索題
    rng = random.Random(37)
    qs = []
    for clue, w in C.CLUE:
        pool = [x for x in allw if x not in group_of(w)]
        wr = rng.sample(pool, 3)
        qs.append(shuffled_mc(clue + "這個詞語是：", w, wr, rng, explain(w)))
    secs.append(("五、本義線索題（依本義提示選出詞語）", "", qs))
    # 六、引申義選擇
    rng = random.Random(51)
    qs = []
    for d in C.W:
        w = d['w']
        g = group_of(w)
        pool = [x for x in C.W if x['w'] not in g and x['s'] != d['s']]
        seen = {d['s']}
        wr = []
        cand = pool[:]
        rng.shuffle(cand)
        for x in cand:
            if x['s'] not in seen:
                seen.add(x['s']); wr.append(x['s'])
            if len(wr) == 3:
                break
        qs.append(shuffled_mc("「%s」的衍生義（引申義）是：" % w, d['s'], wr, rng, explain(w)))
    secs.append(("六、引申義選擇（共 %d 題）" % len(C.W), "", qs))
    return secs


SECS = build_sections()
TOTAL = sum(len(q) for _, _, q in SECS)


def answer_lines():
    out = []
    for title, _, qs in SECS:
        name = title.split('、', 1)[1]
        name = name.split('（')[0]
        if qs[0]['kind'] == 'in':
            line = "　".join("%d.%s" % (i + 1, q['a']) for i, q in enumerate(qs))
        elif qs[0]['kind'] == 'sel':
            line = "　".join("%s－%s" % (q['stem'].replace(' ⇄', ''), q['ans'][2:]) for q in qs)
        elif name in ("詞語運用", "本義線索題"):
            line = "　".join("%d.%s" % (i + 1, q['ans']) for i, q in enumerate(qs))
        else:
            line = "　".join("%d.%s" % (i + 1, q['a']) for i, q in enumerate(qs))
        out.append((name, line))
    return out


# ------------------------------------------------------------------ HTML 互動版
CSS = """
.who{display:flex;flex-wrap:wrap;gap:6px 14px;align-items:center;margin:8px 0}#send-msg{font-size:.9em;font-weight:bold}.xp{font-weight:normal}
:root{--bg:#fbf8f1;--fg:#2b2b2b;--card:#fff;--ac:#2f6f5e;--bd:#e2dccb;--ok:#d9f2e3;--bad:#fbdcdc}
@media screen and (prefers-color-scheme:dark){:root{--bg:#1d1f1e;--fg:#eee;--card:#262928;--ac:#7fd0b6;--bd:#3a3f3d;--ok:#1f4a37;--bad:#5a2b2b}}
body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.7 "Noto Sans TC","PingFang TC","Microsoft JhengHei",sans-serif}
main{max-width:960px;margin:auto;padding:16px}h1{color:var(--ac);margin:.3em 0}h2{border-left:6px solid var(--ac);padding-left:10px;margin-top:2em}
nav a{margin-right:12px;color:var(--ac)}.bar{position:sticky;top:0;background:var(--bg);padding:8px 0;border-bottom:1px solid var(--bd);z-index:5;display:flex;gap:10px;flex-wrap:wrap;align-items:center}
button,select,input{font:inherit}.btn{background:var(--ac);color:#fff;border:0;border-radius:8px;padding:6px 14px;cursor:pointer}
.cat{margin:1.4em 0 .5em;padding:4px 14px;border-radius:8px;color:#fff;font-weight:700}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(290px,1fr));gap:12px}.card{background:var(--card);border:1px solid var(--bd);border-left:6px solid var(--c,var(--ac));border-radius:12px;padding:10px 14px}
.card h3{margin:.2em 0;color:var(--c,var(--ac))}.card h3 span{background:var(--c,var(--ac));color:#fff;border-radius:50%;display:inline-block;width:1.7em;height:1.7em;text-align:center;font-size:.8em;line-height:1.7em;margin-right:8px}
.card p{margin:.3em 0}.k{background:var(--bd);border-radius:5px;padding:0 6px;margin-right:6px;font-size:.85em}.s{background:#cfe8ff;color:#123}.o{background:#ffe0c2;color:#321}.d{background:#e8dcff;color:#213}.dn{font-size:.92em}
.q{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:6px 12px;margin:8px 0}.q p{margin:.3em 0}.opts{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:6px}
.opt{text-align:left;background:transparent;color:var(--fg);border:1px solid var(--bd);border-radius:8px;padding:6px 10px;cursor:pointer}.opt.sel{border-color:var(--ac);background:var(--ok)}
.q.ok{background:var(--ok)}.q.bad{background:var(--bad)}.fb{font-size:.9em;font-weight:bold}details{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:8px 14px}
.dl a{display:inline-block;margin:6px 10px 6px 0;padding:8px 14px;border-radius:8px;background:var(--ac);color:#fff;text-decoration:none}
"""

JS = r"""
const T=__T__,QUIZ="__QUIZ__";
if(!document.querySelector('h3[data-sec]')){let si=-1,el=document.getElementById('test');while(el&&(el=el.nextElementSibling)){if(el.tagName==='H2')break;if(el.tagName==='H3'){si++;el.dataset.sec=si}else if(el.classList.contains('q'))el.dataset.s=si;else el.querySelectorAll('.q').forEach(q=>q.dataset.s=si)}}
const Q=[...document.querySelectorAll('.q')];
const $=id=>document.getElementById(id);
const val=q=>{const ins=[...q.querySelectorAll('input[type=text]')];if(ins.length>1){const v=ins.map(i=>i.value.trim());return v.every(x=>x)?v.join('|'):''}const i=ins[0],s=q.querySelector('select'),o=q.querySelector('.opt.sel');return i?i.value.trim():s?s.value:o?o.dataset.v:''};
const good=(q,v)=>v===q.dataset.a||(!!q.dataset.alt&&(v===q.dataset.alt||q.dataset.alt.split('｜').includes(v)));
function fb(q,t){const f=q.querySelector('.fb');f.textContent=t;if(q.dataset.x&&t){const x=document.createElement('div');x.className='xp';x.textContent='解析：'+q.dataset.x;f.appendChild(x)}}
function mark(q){const v=val(q);if(!v){q.classList.remove('ok','bad');fb(q,'');return}const ok=good(q,v);q.classList.toggle('ok',ok);q.classList.toggle('bad',!ok);fb(q,ok?'✔ 答對了':'✘ 正確答案：'+q.dataset.ans)}
const mode=()=>document.querySelector('[name=m]:checked').value;
function stat(){const d=Q.filter(q=>val(q)).length;const r=Q.filter(q=>q.classList.contains('ok')).length;$('sc').textContent='已作答 '+d+'／'+T+'　答對 '+(mode()=='live'?r:'—')}
function upd(q){if(mode()=='live')mark(q);stat()}
Q.forEach(q=>{q.querySelectorAll('.opt').forEach(b=>b.onclick=()=>{q.querySelectorAll('.opt').forEach(x=>x.classList.remove('sel'));b.classList.add('sel');upd(q)});q.querySelectorAll('input,select').forEach(e=>e.addEventListener(e.tagName=='SELECT'?'change':'input',()=>upd(q)))});
document.querySelectorAll('[name=m]').forEach(r=>r.onchange=()=>{$('go').hidden=mode()=='live';Q.forEach(q=>{q.classList.remove('ok','bad');fb(q,'')});if(mode()=='live')Q.forEach(mark);stat()});
__SEND__
$('go').onclick=()=>{Q.forEach(q=>{if(val(q))mark(q);else{q.classList.add('bad');fb(q,'未作答，正確答案：'+q.dataset.ans)}});const r=Q.filter(q=>q.classList.contains('ok')).length;$('sc').textContent='已作答 '+Q.filter(q=>val(q)).length+'／'+T+'　答對 '+r+'　（'+Math.round(r*100/T)+' 分）';__CALL__};
$('rs').onclick=()=>{Q.forEach(q=>{q.classList.remove('ok','bad');fb(q,'');q.querySelectorAll('.opt').forEach(x=>x.classList.remove('sel'));q.querySelectorAll('input,select').forEach(e=>e.value='')});$('send-msg').textContent='';stat()};
"""

SEND = r"""function send(r){
  const msg=$('send-msg'),cls=$('f-cls').value.trim(),seat=$('f-seat').value.trim();
  if(!window.GuwenAuth||!GuwenAuth.user){msg.textContent='（尚未用學校帳號登入，成績沒有送出）';if(window.GuwenAuth)GuwenAuth.requireLogin();return}
  if(!cls||!seat){msg.textContent='（請填班級、座號後再按一次「交卷送出」）';(!cls?$('f-cls'):$('f-seat')).focus();return}
  const detail={};let ans='';
  document.querySelectorAll('h3[data-sec]').forEach(h=>{const qs=Q.filter(q=>q.dataset.s===h.dataset.sec);detail[h.textContent.replace(/^[一二三四五六七八九十]+、/,'').replace(/（.*?）/g,'')]=qs.filter(q=>q.classList.contains('ok')).length+'/'+qs.length});
  Q.forEach(q=>{ans+=!val(q)?'-':q.classList.contains('ok')?'1':'0'});
  const total=Math.round(r*100/T);msg.textContent='送出中…';
  fetch('/api/scores',{method:'POST',credentials:'same-origin',headers:{'Content-Type':'application/json'},body:JSON.stringify({quiz:QUIZ,cls:cls,seat:seat,mode:'交卷',total:total,detail:detail,answers:ans})})
   .then(x=>{if(x.status===401)throw'login';return x.json()})
   .then(j=>{msg.textContent=j&&j.ok?'✓ 成績已送出給老師（'+total+' 分）':'送出失敗，請再按一次「交卷送出」'})
   .catch(e=>{msg.textContent=e==='login'?'登入已過期，請重新登入後再按一次「交卷送出」':'送出失敗，請檢查網路後再按一次「交卷送出」'});
}"""

NAVS = {
    "review": '<p><a href="../">← 總複習</a>　<a href="static.html">講義版（可列印）</a></p>',
    "book1": '<p><a href="../">← 心囚　課文與字音字形</a>　<a href="static.html">講義版（可列印）</a></p>',
    "standalone": '<p><a href="static.html">講義版（可列印）</a></p>',
}


def card_html(i, d):
    col = C.CATS[d['cat']][1]
    h = '<div class="card" style="--c:%s"><h3><span>%d</span>%s</h3>' % (col, i, esc(d['w']))
    for c, t in d['parts']:
        h += '<p><b class="k s">%s</b>%s</p>' % (esc(c), esc(t))
    h += '<p><b class="k o">衍生義</b>%s「%s」</p>' % (d['kind'], esc(d['d']))
    h += '<p><b class="k">例句</b>%s</p>' % esc(d['ex'])
    if d.get('x'):
        h += '<p class="dn"><b class="k d">補充</b>%s</p>' % esc(d['x'])
    return h + '</div>'


def q_html(n, q):
    attrs = 'data-a="%s" data-ans="%s"' % (esc(q['a']), esc(q['ans']))
    if q.get('x'):
        attrs += ' data-x="%s"' % esc(q['x'])
    if q['kind'] == 'in':
        stem = esc(q['stem']).replace('「', '「').replace('」', '」')
        return '<div class="q" %s><p>%d. %s　<input type="text" maxlength="2" size="3" autocomplete="off"></p><div class="fb"></div></div>' % (attrs, n, stem)
    if q['kind'] == 'sel':
        opts = '<option value="">請選擇</option>' + ''.join('<option value="%s">(%s) %s</option>' % (L[k], L[k], esc(o)) for k, o in enumerate(q['opts']))
        return '<div class="q" %s><p>%d. %s <select>%s</select></p><div class="fb"></div></div>' % (attrs, n, esc(q['stem']), opts)
    opts = ''.join('<button type="button" class="opt" data-v="%s">(%s) %s</button>' % (L[k], L[k], esc(o)) for k, o in enumerate(q['opts']))
    return '<div class="q" %s><p>%d. %s</p><div class="opts">%s</div><div class="fb"></div></div>' % (attrs, n, esc(q['stem']), opts)


def build_index(variant):
    site = variant != "standalone"
    h = '<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>%s｜互動版</title>\n' % TITLE
    if site:
        h += '<script src="/assets/guwen-qr.js" defer></script>\n<script src="/assets/quiz-auth.js" defer></script>\n'
    h += '<style>' + CSS + '</style>\n</head><body><main>\n' + NAVS[variant]
    if site:
        h += '<div id="page-qr" data-title="%s"></div>' % TITLE
    submit_word = "交卷送出" if site else "交卷"
    intro = ("%d 個雙詞學習站、字形測驗、綜合測驗、連連看、詞語運用（含 %d 個例句應用題）、本義線索題與引申義選擇，共 %d 題。直接點選或輸入作答；選「交卷模式」作答完按「%s」%s。"
             % (len(C.W), len(C.W), TOTAL, submit_word, "，成績會送給老師" if site else "，會顯示分數"))
    h += '<h1>%s（中學生）</h1><p>%s</p>\n' % (TITLE, intro)
    h += '<nav><a href="#learn">雙詞學習站</a><a href="#test">實力挑戰</a><a href="#extra">大補帖</a><a href="#answers">參考答案</a><a href="#dl">下載 Word</a></nav>\n'
    if site:
        h += ('<div class="who" id="who"><label>班級 <input id="f-cls" size="5" maxlength="8" placeholder="例：711" autocomplete="off"></label>'
              '<label>座號 <input id="f-seat" size="3" maxlength="2" inputmode="numeric" autocomplete="off"></label>'
              '<label>姓名 <input id="f-name" size="8" maxlength="10" autocomplete="off"></label></div>\n')
    h += ('<div class="bar"><label><input type="radio" name="m" value="live" checked> 即時對答案</label><label><input type="radio" name="m" value="sub"> 交卷模式</label>\n'
          '<span id="sc">已作答 0／%d　答對 0</span><button class="btn" id="go" type="button" hidden>%s</button><button class="btn" id="rs" type="button">重新作答</button><span id="send-msg"></span></div>\n' % (TOTAL, submit_word))
    h += '<h2 id="learn">雙詞學習站（共 %d 個）</h2>\n' % len(C.W)
    n = 0
    for ci, (cname, col) in enumerate(C.CATS):
        h += '<div class="cat" style="background:%s">%s</div><div class="grid">' % (col, cname)
        for d in C.W:
            if d['cat'] == ci:
                n += 1
                h += card_html(n, d)
        h += '</div>\n'
    h += '<h2 id="test">實力挑戰</h2>'
    for title, _, qs in SECS:
        h += '<h3>%s</h3>' % esc(title)
        for i, q in enumerate(qs):
            h += q_html(i + 1, q)
        h += '\n'
    h += '<h2 id="extra">大補帖：同類詞語整理</h2><div class="card">'
    for k, v in C.EXTRA:
        h += '<p><b class="k">%s</b>%s</p>' % (esc(k), esc(v))
    h += '</div>\n<h2 id="answers">參考答案</h2><details><summary>顯示參考答案（老師用）</summary>'
    for name, line in answer_lines():
        h += '<p>%s：%s</p>' % (esc(name), esc(line))
    h += '<p style="font-size:.9em;color:var(--fg)">%s</p></details>\n' % esc(C.NOTE)
    h += ('<h2 id="dl">下載 Word 檔</h2><div class="dl"><a href="shuangci-handout.docx">📄 標準版 Word 檔（黑白，適合列印）</a>'
          '<a href="shuangci-handout-color.docx">🎨 彩色版 Word 檔（內容相同，依分類上色）</a></div>\n</main><script>')
    js = JS.replace('__T__', str(TOTAL)).replace('__QUIZ__', TITLE)
    js = js.replace('__SEND__', SEND if site else '').replace('__CALL__', 'send(r)' if site else '')
    h += js + '\n</script></body></html>\n'
    return h


# ------------------------------------------------------------------ 講義版 HTML
def build_static(variant):
    site = variant != "standalone"
    h = '<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>%s｜講義版</title>\n' % TITLE
    if site:
        h += '<script src="/assets/guwen-qr.js" defer></script>\n'
    h += ('<style>body{font:15px/1.7 "Noto Sans TC","Microsoft JhengHei",sans-serif;color:#000;background:#fff;max-width:820px;margin:auto;padding:16px}h1,h2{color:#1f5c5a}'
          '.c{break-inside:avoid;margin:8px 0;padding:6px 0;border-bottom:1px solid #ccc}.q{break-inside:avoid;margin:6px 0}.o{margin:0 0 4px 1.5em}.pb{break-before:page}.nt{font-size:.9em;color:#444}@media print{.np{display:none}}</style>\n</head><body>\n')
    back = {"review": '<a href="../">← 總複習</a>　', "book1": '<a href="../">← 心囚　課文與字音字形</a>　', "standalone": ''}[variant]
    h += '<p class="np">%s<a href="%s">互動版</a>　<a href="shuangci-handout.docx">Word 講義</a>　<a href="shuangci-handout-color.docx">Word 講義（彩色版）</a></p>\n' % (back, "./" if site else "index.html")
    if site:
        h += '<div id="page-qr" data-title="%s"></div>' % TITLE
    h += '<h1>%s（中學生）</h1><p>班級：＿＿＿＿　座號：＿＿＿＿　姓名：＿＿＿＿＿＿　得分：＿＿＿＿</p>\n' % TITLE
    h += '<h2>壹、雙詞學習站</h2>'
    n = 0
    for ci, (cname, col) in enumerate(C.CATS):
        h += '<h3 style="color:#1f5c5a">%s</h3>' % cname
        for d in C.W:
            if d['cat'] == ci:
                n += 1
                h += '<div class="c"><b>%d %s</b><br>' % (n, esc(d['w']))
                for c, t in d['parts']:
                    h += '%s：%s<br>' % (esc(c), esc(t))
                h += '衍生義：%s「%s」<br>例句：%s' % (d['kind'], esc(d['d']), esc(d['ex']))
                if d.get('x'):
                    h += '<br>補充：%s' % esc(d['x'])
                h += '</div>'
    h += '<h2 class="pb">貳、實力挑戰</h2>'
    for title, _, qs in SECS:
        h += '<h3>%s</h3>' % esc(title)
        if qs[0]['kind'] == 'sel':
            h += '<p class="nt">%s</p>' % esc("　".join("(%s) %s" % (L[k], o) for k, o in enumerate(qs[0]['opts'])))
        for i, q in enumerate(qs):
            if q['kind'] == 'in':
                h += '<div class="q">%d. %s　＿＿＿</div>' % (i + 1, esc(q['stem']))
            elif q['kind'] == 'sel':
                h += '<div class="q">（　）%d. %s</div>' % (i + 1, esc(q['stem']))
            elif max(len(o) for o in q['opts']) <= 8:
                h += '<div class="q">（　）%d. %s<br>%s</div>' % (i + 1, esc(q['stem']), '　'.join('(%s) %s' % (L[k], esc(o)) for k, o in enumerate(q['opts'])))
            else:
                h += '<div class="q">（　）%d. %s' % (i + 1, esc(q['stem'])) + ''.join('<div class="o">(%s) %s</div>' % (L[k], esc(o)) for k, o in enumerate(q['opts'])) + '</div>'
    h += '<h2>參、大補帖：同類詞語整理</h2>'
    for k, v in C.EXTRA:
        h += '<p><b>%s</b>　%s</p>' % (esc(k), esc(v))
    h += '<h2 class="pb">參考答案</h2>'
    for name, line in answer_lines():
        h += '<p>%s：%s</p>' % (esc(name), esc(line))
    h += '<p class="nt">%s</p></body></html>\n' % esc(C.NOTE)
    return h


# ------------------------------------------------------------------ Word
def build_docx(color, path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    FONT = "Microsoft JhengHei"

    def setfont(run, size=None, bold=None, rgb=None):
        run.font.name = FONT
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            rf.set(qn(a), FONT)
        if size: run.font.size = Pt(size)
        if bold is not None: run.font.bold = bold
        if rgb: run.font.color.rgb = RGBColor.from_string(rgb)

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, m, Cm(2.0))
    st = d.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.2
    TEAL = "1F5C5A" if color else "000000"
    for n, sz in (('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)):
        hs = d.styles[n]
        hs.font.name = FONT; hs.font.size = Pt(sz); hs.font.bold = True
        hs.font.color.rgb = RGBColor.from_string(TEAL)
        hs.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        hs.paragraph_format.space_before = Pt(12); hs.paragraph_format.space_after = Pt(4)

    def para(text="", bold=False, size=None, rgb=None, indent=None, align=None):
        p = d.add_paragraph()
        if indent: p.paragraph_format.left_indent = Cm(indent)
        if align: p.alignment = align
        r = p.add_run(text); setfont(r, size, bold, rgb)
        return p

    def shade(cell, fill):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

    def cell_text(cell, text, bold=False, size=10, rgb=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text); setfont(r, size, bold, rgb)

    d.add_heading("%s（中學生）" % TITLE, level=1)
    para("班級：＿＿＿＿　座號：＿＿＿＿　姓名：＿＿＿＿＿＿　得分：＿＿＿＿")
    d.add_heading("壹、雙詞學習站", level=2)
    widths = [Cm(2.0), Cm(5.4), Cm(4.6), Cm(5.0)]
    n = 0
    for ci, (cname, col) in enumerate(C.CATS):
        col = col.lstrip('#')
        d.add_heading(cname, level=3)
        words = [x for x in C.W if x['cat'] == ci]
        t = d.add_table(rows=1, cols=4)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for j, name in enumerate(("詞語", "本義", "衍生義", "例句")):
            c = t.rows[0].cells[j]; c.width = widths[j]
            cell_text(c, name, True, 10, "FFFFFF" if color else None)
            shade(c, col if color else "D9D9D9")
        for x in words:
            n += 1
            row = t.add_row().cells
            for j in range(4): row[j].width = widths[j]
            cell_text(row[0], "%d %s" % (n, x['w']), True, 11, col if color else None)
            row[1].text = ""
            p = row[1].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            for k, (c, tx) in enumerate(x['parts']):
                if k: p.add_run().add_break()
                r = p.add_run(c + "："); setfont(r, 10, True)
                r = p.add_run(tx); setfont(r, 10)
            cell_text(row[2], "%s「%s」" % (x['kind'], x['d']), False, 10)
            ex = x['ex'] + (("\n補充：" + x['x']) if x.get('x') else "")
            row[3].text = ""
            p = row[3].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(x['ex']); setfont(r, 10)
            if x.get('x'):
                p.add_run().add_break()
                r = p.add_run("補充：" + x['x']); setfont(r, 9, False, "555555")
            if color:
                shade(row[0], "F3F6F4")
        para("")
    d.add_page_break()
    d.add_heading("貳、實力挑戰", level=2)
    for title, _, qs in SECS:
        d.add_heading(title, level=3)
        if qs[0]['kind'] == 'sel':
            para("　".join("(%s) %s" % (L[k], o) for k, o in enumerate(qs[0]['opts'])), size=10)
        for i, q in enumerate(qs):
            if q['kind'] == 'in':
                para("%d. %s　＿＿＿" % (i + 1, q['stem']))
            elif q['kind'] == 'sel':
                para("（　）%d. %s" % (i + 1, q['stem']))
            else:
                para("（　）%d. %s" % (i + 1, q['stem']))
                if max(len(o) for o in q['opts']) <= 8:
                    para("　".join("(%s) %s" % (L[k], o) for k, o in enumerate(q['opts'])), indent=1.0)
                else:
                    for k, o in enumerate(q['opts']):
                        para("(%s) %s" % (L[k], o), indent=1.0)
    d.add_heading("參、大補帖：同類詞語整理", level=2)
    for k, v in C.EXTRA:
        p = d.add_paragraph()
        r = p.add_run(k + "　"); setfont(r, 11, True, TEAL)
        r = p.add_run(v); setfont(r, 11)
    d.add_page_break()
    d.add_heading("參考答案", level=2)
    for name, line in answer_lines():
        p = d.add_paragraph()
        r = p.add_run(name + "："); setfont(r, 11, True, TEAL)
        r = p.add_run(line); setfont(r, 11)
    para(C.NOTE, size=9, rgb="555555")
    d.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "out"
    os.makedirs(out, exist_ok=True)
    for v in ("review", "book1", "standalone"):
        od = os.path.join(out, v); os.makedirs(od, exist_ok=True)
        open(os.path.join(od, "index.html"), "w", encoding="utf-8").write(build_index(v))
        open(os.path.join(od, "static.html"), "w", encoding="utf-8").write(build_static(v))
    build_docx(False, os.path.join(out, "shuangci-handout.docx"))
    build_docx(True, os.path.join(out, "shuangci-handout-color.docx"))
    print("TOTAL", TOTAL, [len(q) for _, _, q in SECS])
    for name, line in answer_lines():
        print(name, line[:120])
